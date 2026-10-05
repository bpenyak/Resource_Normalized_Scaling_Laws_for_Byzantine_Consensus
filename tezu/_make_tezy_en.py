# -*- coding: utf-8 -*-
"""English conference theses (max 2 pages), matching instructions.md and the PUTO sample."""
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

OUT = r"C:\pol\paper_3\tezu\Peniak B O_theses.docx"
LANG = "en-US"

NB_HYPHEN = "\u2011"  # rendered as Word non-breaking hyphen


def set_run_font(run, name="Times New Roman", size=11, italic=False, bold=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.italic = italic
    run.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:cs"), name)
    rFonts.set(qn("w:eastAsia"), name)
    lang = rPr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rPr.append(lang)
    lang.set(qn("w:val"), LANG)
    lang.set(qn("w:eastAsia"), LANG)


def add_run(
    p,
    text,
    *,
    italic=False,
    bold=False,
    sub=False,
    sup=False,
    font="Times New Roman",
    size=11,
    all_caps=False,
):
    text = re.sub(r"(?<=\w)-(?=\w)", NB_HYPHEN, text)
    run = p.add_run(text)
    if NB_HYPHEN in text:
        for t in run._element.findall(qn("w:t")):
            run._element.remove(t)
        for i, part in enumerate(text.split(NB_HYPHEN)):
            if i:
                run._element.append(OxmlElement("w:noBreakHyphen"))
            if part:
                t = OxmlElement("w:t")
                t.set(qn("xml:space"), "preserve")
                t.text = part
                run._element.append(t)
    set_run_font(run, name=font, size=size, italic=italic, bold=bold)
    if sub:
        run.font.subscript = True
    if sup:
        run.font.superscript = True
    if all_caps:
        run.font.all_caps = True
    return run


def set_para_format(
    p,
    *,
    align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    first_line=None,
    left=None,
    hanging=None,
    space_before=0,
    space_after=0,
    line=240,
):
    pf = p.paragraph_format
    pf.alignment = align
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pPr = p._p.get_or_add_pPr()
    spacing = pPr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        pPr.append(spacing)
    spacing.set(qn("w:before"), str(int(space_before * 20)))
    spacing.set(qn("w:after"), str(int(space_after * 20)))
    spacing.set(qn("w:line"), str(line))
    spacing.set(qn("w:lineRule"), "auto")
    if first_line is not None:
        pf.first_line_indent = first_line
    else:
        pf.first_line_indent = Cm(0)
    if left is not None:
        pf.left_indent = left
    if hanging is not None:
        pf.first_line_indent = -hanging
        pf.left_indent = hanging


def empty_line(doc):
    p = doc.add_paragraph()
    set_para_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=Cm(0))
    add_run(p, "")
    return p


def add_formula(doc):
    p = doc.add_paragraph()
    set_para_format(
        p,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_line=Cm(0),
        space_before=6,
        space_after=6,
    )
    add_run(p, "TPS(", font="Cambria Math")
    add_run(p, "n", font="Cambria Math", italic=True)
    add_run(p, ", ", font="Cambria Math")
    add_run(p, "c", font="Cambria Math", italic=True)
    add_run(p, ") = ", font="Cambria Math")
    add_run(p, "T", font="Cambria Math", italic=True)
    add_run(p, "0", font="Cambria Math", sub=True)
    add_run(p, " · ", font="Cambria Math")
    add_run(p, "g", font="Cambria Math", italic=True)
    add_run(p, "(", font="Cambria Math")
    add_run(p, "c", font="Cambria Math", italic=True)
    add_run(p, ") · ", font="Cambria Math")
    add_run(p, "n", font="Cambria Math", italic=True)
    add_run(p, "−β", font="Cambria Math", sup=True)
    add_run(p, ",     ", font="Cambria Math")
    add_run(p, "g", font="Cambria Math", italic=True)
    add_run(p, "(", font="Cambria Math")
    add_run(p, "c", font="Cambria Math", italic=True)
    add_run(p, ") = ", font="Cambria Math")
    add_run(p, "c", font="Cambria Math", italic=True)
    add_run(p, "γ", font="Cambria Math", sup=True)
    add_run(p, " / [1 + (", font="Cambria Math")
    add_run(p, "c", font="Cambria Math", italic=True)
    add_run(p, "/", font="Cambria Math")
    add_run(p, "c", font="Cambria Math", italic=True)
    add_run(p, "*", font="Cambria Math", sup=True)
    add_run(p, ")", font="Cambria Math")
    add_run(p, "θ", font="Cambria Math", sup=True)
    add_run(p, "].", font="Cambria Math")


def main():
    doc = Document()

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(2.3))
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), "Times New Roman")
    rFonts.set(qn("w:hAnsi"), "Times New Roman")
    rFonts.set(qn("w:cs"), "Times New Roman")

    p = doc.add_paragraph()
    set_para_format(p, align=WD_ALIGN_PARAGRAPH.LEFT, first_line=Cm(0))
    add_run(p, "UDC: 004.75:519.2")

    p = doc.add_paragraph()
    set_para_format(p, align=WD_ALIGN_PARAGRAPH.RIGHT, first_line=Cm(0))
    add_run(p, "PENIAK B. O., LIUBINSKIY B. B.", italic=True, bold=True, all_caps=True)

    p = doc.add_paragraph()
    set_para_format(p, align=WD_ALIGN_PARAGRAPH.RIGHT, first_line=Cm(0))
    add_run(p, "Lviv Polytechnic National University (Ukraine)")

    empty_line(doc)

    p = doc.add_paragraph()
    set_para_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=Cm(0))
    add_run(
        p,
        "RESOURCE-NORMALIZED SCALING LAWS FOR BFT CONSENSUS",
        bold=True,
        all_caps=True,
    )

    empty_line(doc)

    p = doc.add_paragraph()
    set_para_format(p, first_line=Cm(0.9))
    add_run(
        p,
        "Sizing a permissioned BFT validator set for e‑governance trades falling "
        "throughput against rising fault-detection power as ",
        italic=True,
    )
    add_run(p, "n", italic=True)
    add_run(
        p,
        " grows. We propose resource-normalized measurement and a bi-factor model; "
        "in the unsaturated regime with zero-mean noise the single-factor exponent "
        "tends to ",
        italic=True,
    )
    add_run(p, "γλ", italic=True)
    add_run(p, " − ", italic=True)
    add_run(p, "β", italic=True)
    add_run(
        p,
        " (an omitted-variable bias). On commodity CI for Besu/QBFT at a fixed "
        "quota ",
        italic=True,
    )
    add_run(p, "β", italic=True)
    add_run(
        p,
        " = 1.651 ± 0.151; with a near-chance detector the detection constraint is "
        "relaxed (not enforced) and throughput alone sets the window [4, 4].",
        italic=True,
    )

    empty_line(doc)

    p = doc.add_paragraph()
    set_para_format(p, first_line=Cm(0.9))
    add_run(
        p,
        "Problem statement. Permissioned Byzantine fault-tolerant (BFT) consensus "
        "underpins e‑governance systems and poses a procurement question: how many "
        "validators ",
    )
    add_run(p, "n", italic=True)
    add_run(p, " should the network have? Throughput favours small ")
    add_run(p, "n", italic=True)
    add_run(
        p,
        ": every validator executes every transaction, and PBFT-family protocols "
        "exchange ",
    )
    add_run(p, "O", italic=True)
    add_run(p, "(")
    add_run(p, "n", italic=True)
    add_run(p, "2", sup=True)
    add_run(p, ") messages per block [1]. Fault-detection reliability favours large ")
    add_run(p, "n", italic=True)
    add_run(
        p,
        ". At the design stage, dedicated multi-host testbeds are often unavailable, "
        "while commodity CI runners co-locate many validators on few cores. Without "
        "normalization, single-factor fits TPS(",
    )
    add_run(p, "n", italic=True)
    add_run(p, ") ∝ ")
    add_run(p, "n", italic=True)
    add_run(p, "α", sup=True)
    add_run(
        p,
        " confound consensus cost with host-resource contention and with the load "
        "generator’s concurrency schedule. Published exponents then diverge in "
        "magnitude and sometimes in sign; previous calibrations, extrapolated to "
        "sixty validators, differ by about an order of magnitude [2, 3].",
    )

    p = doc.add_paragraph()
    set_para_format(p, first_line=Cm(0.9))
    add_run(
        p,
        "The following tasks are addressed: separate the algorithmic cost of "
        "consensus from host-resource contention; explain the non-identifiability "
        "of the single-factor exponent along a concurrency path; formulate a "
        "chance-constrained validator-set sizing problem with throughput and "
        "detection constraints; calibrate the workflow on a reproducible "
        "commodity-CI campaign.",
    )

    p = doc.add_paragraph()
    set_para_format(p, first_line=Cm(0.9))
    add_run(
        p,
        "Essence of the study. A resource-normalized measurement (RNM) protocol is "
        "introduced: each validator container receives a fixed CPU quota ",
    )
    add_run(p, "q", italic=True)
    add_run(p, " independent of ")
    add_run(p, "n", italic=True)
    add_run(p, ", and normalized throughput equals TPS / ")
    add_run(p, "q", italic=True)
    add_run(p, ". Under RNM the bi-factor scaling characterization takes the form")

    add_formula(doc)

    p = doc.add_paragraph()
    set_para_format(p, first_line=Cm(0))
    add_run(p, "where ")
    add_run(p, "c", italic=True)
    add_run(
        p,
        " is client concurrency, β > 0 is the replica-count decay exponent, "
        "γ ∈ (0, 1] is the concurrency exponent, and ",
    )
    add_run(p, "c", italic=True)
    add_run(p, "*", italic=True, sup=True)
    add_run(p, " is the saturation threshold (a USL-style form [4]). In the unsaturated regime ")
    add_run(p, "c", italic=True)
    add_run(p, " ≪ ")
    add_run(p, "c", italic=True)
    add_run(p, "*", italic=True, sup=True)
    add_run(p, " the relation is log-linear. If measurements follow the path log ")
    add_run(p, "c", italic=True)
    add_run(p, " = λ log ")
    add_run(p, "n", italic=True)
    add_run(
        p,
        " + μ and the noise has zero mean and finite variance, the OLS slope of "
        "the single-factor model satisfies α̂ →ᴾ γλ − β; this is the standard "
        "omitted-variable bias [5]. In particular α̂ = −β only under a factorial "
        "design λ = 0; without reporting λ, published single-factor exponents "
        "should not be treated as transferable protocol properties across testbeds.",
    )

    p = doc.add_paragraph()
    set_para_format(p, first_line=Cm(0.9))
    add_run(
        p,
        "The calibrated exponent is the input to a chance-constrained program [6]: "
        "minimize a nondecreasing Cost(",
    )
    add_run(p, "n", italic=True)
    add_run(
        p,
        ") subject to the lower prediction-interval bound on throughput meeting "
        "peak demand and the probability of missing a coordinated "
        "transaction-omission fault (an engineering Hoeffding-type bound [7]) not "
        "exceeding a prescribed level. Admissible sizes are discrete (efficient "
        "BFT sets have the form 3f + 1), and the feasible set is a closed interval [",
    )
    add_run(p, "n", italic=True)
    add_run(p, "min", italic=True, sub=True)
    add_run(p, ", ")
    add_run(p, "n", italic=True)
    add_run(p, "max", italic=True, sub=True)
    add_run(
        p,
        "]; an empty interval certifies that no validator-set size meets both "
        "requirements, so the architecture or hardware class must change, not "
        "merely ",
    )
    add_run(p, "n", italic=True)
    add_run(p, ".")

    p = doc.add_paragraph()
    set_para_format(p, first_line=Cm(0.9))
    add_run(
        p,
        "The campaign ran in a public CI workflow on a four-core host: a "
        "Hyperledger Besu/QBFT network under the RNM protocol, 94 successful "
        "records, ",
    )
    add_run(p, "n", italic=True)
    add_run(p, " ≤ 16. Joint OLS at the reference quota ")
    add_run(p, "q", italic=True)
    add_run(p, " = 0.25 yields ")
    add_run(p, "β", italic=True)
    add_run(p, " = 1.651 ± 0.151 and ")
    add_run(p, "γ", italic=True)
    add_run(p, " = 0.178 ± 0.070 (")
    add_run(p, "R", italic=True)
    add_run(p, "² = 0.727). Resampling along ")
    add_run(p, "λ", italic=True)
    add_run(p, " = 1 gives ")
    add_run(p, "α̂", italic=True)
    add_run(p, " = −1.362 versus the analytic ")
    add_run(p, "γλ", italic=True)
    add_run(p, " − ")
    add_run(p, "β", italic=True)
    add_run(p, " = −1.473. Invariance of ")
    add_run(p, "β", italic=True)
    add_run(p, " across ")
    add_run(p, "q", italic=True)
    add_run(p, " ∈ {0.20, 0.25, 0.33} is rejected (")
    add_run(p, "p", italic=True)
    add_run(p, " = 0.002): at ")
    add_run(p, "q", italic=True)
    add_run(
        p,
        " = 0.33 throughput no longer grows with the quota and the total "
        "reservation exceeds host capacity. RNM therefore stabilizes ",
    )
    add_run(p, "β", italic=True)
    add_run(p, " at a fixed quota, and coefficients are reported at ")
    add_run(p, "q", italic=True)
    add_run(
        p,
        " = 0.25. The omission/duplication detector has AUC = 0.506 (near chance); "
        "with it the detection constraint is unsatisfiable (it would require ",
    )
    add_run(p, "n", italic=True)
    add_run(p, "min", italic=True, sub=True)
    add_run(
        p,
        " ≈ 3.3·10³), so the formal feasible set is empty; the constraint is "
        "relaxed (not enforced) and throughput alone sets the window. The headline "
        "fit uses all ",
    )
    add_run(p, "n", italic=True)
    add_run(p, "; a separate refit on ")
    add_run(p, "n", italic=True)
    add_run(p, " ≤ 9 only (")
    add_run(p, "β", italic=True)
    add_run(p, " = 1.300) gives intervals covering 20 and 21 of 22 hold-out runs (90.9% / 95.5%) at ")
    add_run(p, "n", italic=True)
    add_run(p, " ∈ {11, 13, 16}. The constant ")
    add_run(p, "β", italic=True)
    add_run(p, " averages a slope that steepens with ")
    add_run(p, "n", italic=True)
    add_run(p, " (local values ≈ 1.4 → 2.8), which we attribute to residual host contention as the reservation ")
    add_run(p, "nq", italic=True)
    add_run(p, " approaches the four host cores (RNM removes contention only while ")
    add_run(p, "nq", italic=True)
    add_run(p, " < host cores). The full-range ")
    add_run(p, "β", italic=True)
    add_run(p, " is conservative for ")
    add_run(p, "n", italic=True)
    add_run(p, "max", italic=True, sub=True)
    add_run(
        p,
        " but optimistic for κ (13.0 vs. 16.7 under the small-n fit). A scenario of "
        "≈ 3·10⁷ ballots in 12 h gives a mean of "
        "≈ 694 TPS; with an assumed peak-to-mean factor of ≈ 3 the peak is "
        "≈ 2.1·10³ TPS. Fault-free CI runs commit at most ≈ 89 TPS (at ",
    )
    add_run(p, "n", italic=True)
    add_run(p, " = 7, ")
    add_run(p, "c", italic=True)
    add_run(p, " = 16, ")
    add_run(p, "q", italic=True)
    add_run(p, " = 0.25), so the model works in units of TPS/")
    add_run(p, "q", italic=True)
    add_run(
        p,
        ", and the scale κ = 13.0 is read as cores per validator on the target "
        "hardware. Linearity of throughput in the quota is the key assumption and "
        "the main extrapolation risk: κ is ≈ 52 times the calibrated quota, "
        "whereas throughput already stops growing at ",
    )
    add_run(p, "q", italic=True)
    add_run(
        p,
        " = 0.33; κ is a requirement to verify, not a measured capability. The "
        "window then collapses to [4, 4]. Doubled demand or a 200 ms emulated WAN "
        "RTT (",
    )
    add_run(p, "β", italic=True)
    add_run(
        p,
        " rises from 1.459 to 2.708 on the WAN-experiment grid) at fixed κ empties "
        "the feasible set; κ ≈ 26 or ≈ 73, respectively, would be required.",
    )

    p = doc.add_paragraph()
    set_para_format(p, first_line=Cm(0.9))
    add_run(
        p,
        "Conclusions. Divergence of single-factor BFT scaling exponents can arise "
        "from the measurement design (the concurrency path ",
    )
    add_run(p, "λ", italic=True)
    add_run(
        p,
        "), not only from protocol properties. Resource-normalized calibration "
        "at ",
    )
    add_run(p, "λ", italic=True)
    add_run(
        p,
        " = 0 stabilizes the consensus exponent at a fixed quota and turns it into "
        "a reproducible validator-set sizing workflow; the throughput–detection "
        "trade-off is formulated but not demonstrated with the measured detector. "
        "For the studied Besu/QBFT deployment on commodity CI, "
        "the procurement answer is the smallest admissible set ",
    )
    add_run(p, "n", italic=True)
    add_run(p, " = 4; larger ")
    add_run(p, "n", italic=True)
    add_run(
        p,
        " degrades predicted throughput. Numerical coefficients are "
        "deployment-specific; the methodology is intended for recalibration under "
        "matching measurement conditions. Code, data and CI-campaign artifacts: "
        "https://github.com/bpenyak/Resource_Normalized_Scaling_Laws_for_Byzantine_Consensus.",
    )

    empty_line(doc)

    p = doc.add_paragraph()
    set_para_format(
        p,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_line=Cm(0),
        space_after=6,
    )
    add_run(p, "REFERENCES")

    refs = [
        "Castro M., Liskov B. Practical Byzantine fault tolerance. Proceedings of the Third Symposium on Operating Systems Design and Implementation (OSDI). 1999. P. 173–186.",
        "Peniak B. O., Liubinskiy B. B., Solomka I. R. Minimal-sample performance prediction for Byzantine consensus: the α-calibration method. Mathematical Modeling and Computing. 2026. (accepted for publication).",
        "Peniak B. O., Liubinskiy B. B. Adaptive hybrid consensus mechanism for blockchain-based e‑governance: a machine learning approach. Scientific Bulletin of Uzhhorod University. Series of Mathematics and Informatics. 2026. Vol. 49, No. 2. P. 255–261.",
        "Gunther N. J. Guerrilla capacity planning: a tactical approach to planning for highly scalable applications and services. Berlin: Springer, 2007.",
        "Wooldridge J. M. Econometric analysis of cross section and panel data. 2nd ed. Cambridge, MA: MIT Press, 2010. 1064 p.",
        "Charnes A., Cooper W. W. Chance-constrained programming. Management Science. 1959. Vol. 6, No. 1. P. 73–79.",
        "Hoeffding W. Probability inequalities for sums of bounded random variables. Journal of the American Statistical Association. 1963. Vol. 58, No. 301. P. 13–30.",
    ]
    hang = Cm(0.45)
    for i, text in enumerate(refs, 1):
        p = doc.add_paragraph()
        set_para_format(p, hanging=hang, space_after=3)
        add_run(p, f"{i}.\u00a0{text}")

    doc.save(OUT)
    Path(r"C:\pol\paper_3\tezu\_wrote_en.txt").write_text("ok\n", encoding="utf-8")


if __name__ == "__main__":
    main()
